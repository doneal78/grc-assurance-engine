"""Read a vendor intake file into the public questionnaire form.

CSV and .xlsx use the same column names as the public intake form.
A ServiceNow-shaped JSON or CSV file is a simulated ticket export: common
ticket fields are kept as context, and the custom fields and attachment
metadata map into that same form. Nothing here calls ServiceNow, scores a
vendor, or writes to the database. Confirming an import runs the parsed
form through the existing intake builder.
"""

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from app.services.intake_builder import (
    IntakeError, build_intake, clause_evidence_field, field_choices, run_intake,
)
from app.services.intake_questionnaire import load_questionnaire
from app.services.org_policy import load_org_policy

IMPORT_MAX_BYTES = 256 * 1024
MAX_DATA_ROWS = 20
MAX_SCAN_ROWS = 200
MAX_COLUMNS = 40
MAX_CELL_LENGTH = 500
MAX_ATTACHMENTS = 30
SERVICENOW_KIND = "servicenow_vendor_intake"
SERVICENOW_CATEGORY = "vendor_risk"

FORMAT_LABELS = {
    "questionnaire_csv": "questionnaire CSV",
    "questionnaire_xlsx": "questionnaire workbook",
    "servicenow_json": "simulated ServiceNow JSON",
    "servicenow_csv": "simulated ServiceNow CSV",
    "unrecognized": "file",
}

OPTIONAL_CONTEXT = (
    "sys_id", "short_description", "description", "state", "priority",
    "assignment_group", "opened_at",
)
ATTACHMENT_REQUIRED = ("file_name", "u_intake_field", "u_evidence_option")
ATTACHMENT_OPTIONAL = ("sys_id", "content_type", "size_bytes", "table_name", "table_sys_id")
ROOT_KEYS = ("import_kind", "simulation", "note", "result")


class ImportRejected(ValueError):
    """Raised when a confirmation document is not safe to assess."""

    def __init__(self, preview):
        self.preview = preview
        super().__init__("Import was not applied.")


class _DuplicateKey(Exception):
    def __init__(self, key):
        self.key = key


def questionnaire_columns():
    """Column names, in order, matching the public intake form."""
    questionnaire = load_questionnaire()
    policy = load_org_policy()
    columns = ["vendor_name"]
    columns.extend(spec["field"] for spec in questionnaire["context_questions"])
    columns.extend(spec["field"] for spec in policy["intake_fields"])
    columns.append(questionnaire["access_evidence"]["field"])
    columns.append(questionnaire["trust_center"]["url_field"])
    columns.extend(pointer["field"] for pointer in questionnaire["trust_center"]["pointers"])
    columns.extend(clause_evidence_field(clause["clause_id"])
                   for clause in policy["required_contract_clauses"])
    return columns


def evidence_column_names():
    """Fields whose answer is an evidence option, not a scoring answer."""
    questionnaire = load_questionnaire()
    names = [questionnaire["access_evidence"]["field"]]
    names.extend(pointer["field"] for pointer in questionnaire["trust_center"]["pointers"])
    names.extend(clause_evidence_field(clause["clause_id"])
                 for clause in load_org_policy()["required_contract_clauses"])
    return names


def ticket_body_columns():
    """Questionnaire fields that sit on the ticket rather than an attachment."""
    evidence = set(evidence_column_names())
    return [column for column in questionnaire_columns() if column not in evidence]


def column_guide():
    """Allowed values for the upload page, read from the questionnaire and rules."""
    choices = field_choices()
    questionnaire = load_questionnaire()
    evidence = ", ".join(questionnaire["trust_center"]["evidence_options"])
    url_field = questionnaire["trust_center"]["url_field"]
    guide = []
    for column in questionnaire_columns():
        if column == "vendor_name":
            accepted = "Required text, 80 characters or fewer."
        elif column == url_field:
            accepted = "Blank, or one http:// or https:// URL. The file is not fetched."
        elif column in choices:
            accepted = ", ".join(choices[column])
            if column == "model_training_on_customer_data":
                accepted += ". May be blank only when ai_use is none."
        else:
            accepted = evidence
        guide.append({"column": column, "accepted": accepted})
    return guide


def can_import(preview):
    if preview["problems"] or not preview["rows"]:
        return False
    return all(not row["problems"] and not row["conflicts"] for row in preview["rows"])


def confirmation_document(preview):
    """The previewed rows, and nothing else, for the confirm step."""
    return {
        "format": preview["format"],
        "rows": [
            {"row_number": row["row_number"], "ticket_number": row["ticket_number"], "form": row["form"]}
            for row in preview["rows"]
        ],
    }


def _preview(fmt, filename, problems, rows=None, blank_rows_skipped=0):
    return {
        "format": fmt or "unrecognized",
        "filename": filename,
        "problems": list(problems),
        "blank_rows_skipped": blank_rows_skipped,
        "rows": rows or [],
    }


def _safe_filename(filename):
    base = (filename or "").replace("\\", "/").split("/")[-1].strip()
    if not base or base in {".", ".."}:
        return "upload"
    return base[:120]


def _canonical_map(columns):
    return {column.lower(): column for column in columns}


def _empty_form():
    return {column: "" for column in questionnaire_columns()}


def _type_name(value):
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "a boolean"
    if isinstance(value, (int, float)):
        return "a number"
    if isinstance(value, list):
        return "a list"
    if isinstance(value, dict):
        return "an object"
    return type(value).__name__


def _cell_text(value):
    """Return (text, problem). Dates and numbers are not coerced into answers."""
    if value is None:
        return "", None
    if isinstance(value, str):
        text = value.strip()
        if text[:1] in {"=", "+", "@"} or (text.startswith("\t") and text):
            return "", "spreadsheet formulas are not accepted"
        if len(text) > MAX_CELL_LENGTH:
            return "", f"value is longer than {MAX_CELL_LENGTH} characters"
        return text, None
    if isinstance(value, bool):
        return "", "expected text, got a boolean"
    if isinstance(value, (int, float)):
        return "", "expected text, got a number"
    if isinstance(value, (datetime, date, time)):
        return "", "expected text, got a date"
    return "", f"expected text, got {_type_name(value)}"


def _json_text(value, label, problems):
    if isinstance(value, str):
        text, problem = _cell_text(value)
        if problem:
            problems.append(f"{label}: {problem}.")
            return ""
        return text
    problems.append(f"{label}: expected text, got {_type_name(value)}.")
    return ""


def _normalize_header(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        return None
    return "_".join(value.strip().lower().split())


def _is_blank_row(values):
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return False
    return True


def _trim_width(grid):
    width = 0
    for row in grid:
        for index, value in enumerate(row, start=1):
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            width = max(width, index)
    return [list(row[:width]) for row in grid]


def _blank_row(row_number, form=None):
    return {
        "row_number": row_number,
        "vendor_name": "",
        "ticket_number": None,
        "ticket_context": {},
        "attachments": [],
        "form": form if form is not None else _empty_form(),
        "problems": [],
        "conflicts": [],
    }


def _remember_intake(row, form, as_of, problems):
    if problems:
        row["problems"].extend(problems)
        row["vendor_name"] = form.get("vendor_name", "") if isinstance(form.get("vendor_name"), str) else ""
        row["form"] = form
        return
    try:
        build_intake(form, as_of)
    except IntakeError as exc:
        row["problems"].extend(f"Row {row['row_number']}: {problem}" for problem in exc.problems)
    row["form"] = form
    row["vendor_name"] = form.get("vendor_name", "")


def _apply_conflicts(rows):
    groups = {}
    for row in rows:
        name = row["form"].get("vendor_name", "")
        if not isinstance(name, str) or not name.strip():
            continue
        groups.setdefault(name.strip().casefold(), []).append(row)
    fields = questionnaire_columns()
    for group in groups.values():
        if len(group) < 2:
            continue
        differed = []
        for field in fields:
            values = []
            for row in group:
                value = row["form"].get(field, "")
                values.append(value.strip() if isinstance(value, str) else "")
            if len(set(values)) > 1:
                differed.append((field, values))
        if len(group) == 2:
            numbers = f"{group[0]['row_number']} and {group[1]['row_number']}"
        else:
            numbers = ", ".join(str(row["row_number"]) for row in group)
        label = " and ".join(dict.fromkeys(row["vendor_name"] for row in group))
        if differed:
            detail = "; ".join(
                f"{field} ({' vs '.join(dict.fromkeys(value or 'blank' for value in values))})"
                for field, values in differed
            )
            message = f"Rows {numbers} describe {label} but disagree: {detail}."
        else:
            message = (f"Rows {numbers} describe {label} with the same answers. "
                       "Duplicate rows are kept visible and are not imported.")
        for row in group:
            row["conflicts"].append(message)


def _finish(fmt, filename, file_problems, rows, blank_rows_skipped):
    _apply_conflicts(rows)
    return _preview(fmt, filename, file_problems, rows, blank_rows_skipped)


def _too_many(count):
    return (f"The file has {count} vendor rows. The limit is {MAX_DATA_ROWS}, "
            "so nothing from this file was imported.")


def _classify_table(headers, suffix):
    names = {header for header in headers if header}
    if "u_vendor_name" in names and "vendor_name" in names:
        return None, "The file mixes questionnaire columns and ServiceNow columns."
    servicenow = "u_vendor_name" in names or (any(name.startswith("u_") for name in names) and "vendor_name" not in names)
    if suffix == ".xlsx" and servicenow:
        return None, "Excel import uses the questionnaire columns. A ServiceNow ticket file should be JSON or CSV."
    if suffix == ".xlsx":
        return "questionnaire_xlsx", None
    if servicenow:
        return "servicenow_csv", None
    return "questionnaire_csv", None


def _required_headers(fmt):
    if fmt == "servicenow_csv":
        return ["number", "category", *[f"u_{column}".lower() for column in questionnaire_columns()]]
    return [column.lower() for column in questionnaire_columns()]


def _optional_headers(fmt):
    if fmt == "servicenow_csv":
        return set(OPTIONAL_CONTEXT)
    return set()


def _parse_grid(grid, filename, suffix, as_of, file_problems):
    grid = _trim_width(grid)
    if not grid or _is_blank_row(grid[0]):
        file_problems.append("The file has no header row.")
        return _preview(None, filename, file_problems)
    headers = []
    for index, raw in enumerate(grid[0], start=1):
        name = _normalize_header(raw)
        if name is None:
            file_problems.append(f"Column {index} header must be text.")
            headers.append("")
            continue
        if name[:1] in {"=", "+", "@"}:
            file_problems.append(f"Column {index} header contains a spreadsheet formula.")
        headers.append(name)
    seen = {}
    for index, name in enumerate(headers, start=1):
        if not name:
            continue
        if name in seen:
            file_problems.append(f"Column {name} is duplicated.")
        seen[name] = index
    fmt, problem = _classify_table(headers, suffix)
    if problem:
        file_problems.append(problem)
        return _preview(None, filename, file_problems)
    required = _required_headers(fmt)
    optional = _optional_headers(fmt)
    allowed = set(required) | optional
    missing = [name for name in required if name not in seen]
    unknown = [name for name in headers if name and name not in allowed]
    if missing:
        file_problems.append("Missing columns: " + ", ".join(missing) + ".")
    if unknown:
        file_problems.append("Unrecognized columns: " + ", ".join(unknown) + ". They were not dropped.")
    if any(not name for name in headers):
        file_problems.append("Every used column needs a header.")
    data = []
    blank = 0
    for row_number, row in enumerate(grid[1:], start=2):
        if _is_blank_row(row):
            blank += 1
            continue
        data.append((row_number, row))
    if not data and not missing and not unknown:
        file_problems.append("The file has no vendor rows.")
    if len(data) > MAX_DATA_ROWS:
        file_problems.append(_too_many(len(data)))
        return _preview(fmt, filename, file_problems, blank_rows_skipped=blank)
    if missing or unknown or any(not name for name in headers):
        return _preview(fmt, filename, file_problems, blank_rows_skipped=blank)

    canonical = _canonical_map(questionnaire_columns())
    rows = []
    for row_number, raw_row in data:
        row = _blank_row(row_number)
        mapped = {}
        shape_problems = []
        for index, name in enumerate(headers):
            if not name:
                continue
            raw = raw_row[index] if index < len(raw_row) else None
            text, problem = _cell_text(raw)
            if problem:
                shape_problems.append(f"Row {row_number}, {name}: {problem}.")
            mapped[name] = text
        form = _empty_form()
        if fmt == "servicenow_csv":
            for column in questionnaire_columns():
                form[column] = mapped.get(f"u_{column}".lower(), "")
            number = mapped.get("number", "")
            if not number:
                shape_problems.append(f"Row {row_number}: ticket number is required.")
            else:
                row["ticket_number"] = number
            if mapped.get("category", "") != SERVICENOW_CATEGORY:
                shape_problems.append(
                    f"Row {row_number}, category: expected {SERVICENOW_CATEGORY}.")
            row["ticket_context"] = {key: mapped[key] for key in OPTIONAL_CONTEXT if key in mapped and mapped[key]}
        else:
            for column in questionnaire_columns():
                form[column] = mapped.get(column.lower(), "")
        if shape_problems:
            row["problems"].extend(shape_problems)
            row["form"] = form
            row["vendor_name"] = form.get("vendor_name", "")
        else:
            _remember_intake(row, form, as_of, [])
        rows.append(row)
    return _finish(fmt, filename, file_problems, rows, blank)


def _read_text(data, filename):
    try:
        return data.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        return None, _preview(None, filename, ["The file is not UTF-8 text."])


def _parse_csv(data, filename, as_of):
    text, failed = _read_text(data, filename)
    if failed:
        return failed
    if "\x00" in text:
        return _preview(None, filename, ["The file is not a CSV."], )
    if not text.strip():
        return _preview(None, filename, ["The file is empty."])
    first = text.splitlines()[0]
    if ";" in first and "," not in first:
        return _preview(None, filename, ["CSV files must use commas."])
    try:
        grid = list(csv.reader(io.StringIO(text, newline="")))
    except csv.Error:
        return _preview(None, filename, ["The CSV could not be read."])
    if len(grid) > MAX_SCAN_ROWS:
        return _preview(None, filename, [f"The CSV is longer than {MAX_SCAN_ROWS} records."])
    return _parse_grid(grid, filename, ".csv", as_of, [])


def _sheet_has_values(worksheet):
    limit_row = min(worksheet.max_row or 1, MAX_SCAN_ROWS)
    limit_col = min(worksheet.max_column or 1, MAX_COLUMNS)
    for row in worksheet.iter_rows(max_row=limit_row, max_col=limit_col, values_only=True):
        if not _is_blank_row(row):
            return True
    return False


def _sheet_has_comments(worksheet):
    limit_row = min(worksheet.max_row or 1, MAX_SCAN_ROWS)
    limit_col = min(worksheet.max_column or 1, MAX_COLUMNS)
    for row in worksheet.iter_rows(max_row=limit_row, max_col=limit_col):
        for cell in row:
            if cell.comment is not None:
                return True
    return False


def _parse_xlsx(data, filename, as_of):
    if data[:2] != b"PK":
        return _preview(None, filename, ["This file is not a readable .xlsx workbook."])
    try:
        from openpyxl import load_workbook
    except ImportError:
        return _preview(None, filename, ["Excel import needs openpyxl, which is not installed."])
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        return _preview(None, filename, ["This file is not a readable .xlsx workbook."])
    file_problems = []
    try:
        if workbook.vba_archive is not None:
            file_problems.append("Workbooks with macros are not accepted.")
        if not workbook.sheetnames:
            file_problems.append("The workbook has no worksheets.")
            return _preview("questionnaire_xlsx", filename, file_problems)
        extras = [name for name in workbook.sheetnames[1:] if _sheet_has_values(workbook[name])]
        if extras:
            file_problems.append(
                "Only the first worksheet is imported. These worksheets also have values: "
                + ", ".join(extras) + ".")
        sheet = workbook[workbook.sheetnames[0]]
        if sheet.merged_cells.ranges:
            file_problems.append("Merged cells are not accepted.")
        if _sheet_has_comments(sheet):
            file_problems.append("Cell comments are not imported. Remove them so nothing is left out.")
        if (sheet.max_column or 0) > MAX_COLUMNS:
            file_problems.append(f"The worksheet has more than {MAX_COLUMNS} columns.")
            return _preview("questionnaire_xlsx", filename, file_problems)
        if (sheet.max_row or 0) > MAX_SCAN_ROWS:
            file_problems.append(f"The worksheet is longer than {MAX_SCAN_ROWS} rows.")
            return _preview("questionnaire_xlsx", filename, file_problems)
        grid = [
            list(row) for row in sheet.iter_rows(
                max_row=sheet.max_row or 1, max_col=sheet.max_column or 1, values_only=True)
        ]
    finally:
        workbook.close()
    parsed = _parse_grid(grid, filename, ".xlsx", as_of, file_problems)
    return parsed


def _pairs(pairs):
    found = {}
    for key, value in pairs:
        if key in found:
            raise _DuplicateKey(key)
        found[key] = value
    return found


def _reject_known_export(document, filename):
    if not isinstance(document, dict):
        return None
    if document.get("kind") == "vendor_intake_package" or "export_version" in document:
        return _preview(None, filename, [
            "This file is a reviewer package export. It is not an intake file. "
            "Use the questionnaire columns or a simulated ServiceNow ticket file."])
    if "_demo_note" in document and "result" not in document:
        return _preview(None, filename, [
            "This file is a mock remediation ticket. It is not an intake file, "
            "and it was not sent to ServiceNow."])
    return None


def _attachment_preview(file_name, field, option, content_type, size_bytes, sys_id):
    return {
        "file_name": file_name,
        "intake_field": field,
        "evidence_option": option,
        "content_type": content_type,
        "size_bytes": size_bytes,
        "sys_id": sys_id,
    }


def _parse_attachment(attachment, label, evidence_fields, problems):
    if not isinstance(attachment, dict):
        problems.append(f"{label}: expected an object.")
        return None
    unknown = sorted(set(attachment) - set(ATTACHMENT_REQUIRED) - set(ATTACHMENT_OPTIONAL))
    if unknown:
        problems.append(f"{label}: unrecognized fields {', '.join(unknown)}. They were not dropped.")
    missing = [key for key in ATTACHMENT_REQUIRED if key not in attachment]
    if missing:
        problems.append(f"{label}: missing {', '.join(missing)}.")
    file_name = _json_text(attachment.get("file_name", ""), f"{label} file_name", problems) if "file_name" in attachment else ""
    raw_field = _json_text(attachment.get("u_intake_field", ""), f"{label} u_intake_field", problems) if "u_intake_field" in attachment else ""
    option = _json_text(attachment.get("u_evidence_option", ""), f"{label} u_evidence_option", problems) if "u_evidence_option" in attachment else ""
    content_type = ""
    if "content_type" in attachment:
        content_type = _json_text(attachment.get("content_type"), f"{label} content_type", problems)
    sys_id = ""
    if "sys_id" in attachment:
        sys_id = _json_text(attachment.get("sys_id"), f"{label} sys_id", problems)
    for key in ("table_name", "table_sys_id"):
        if key in attachment:
            _json_text(attachment.get(key), f"{label} {key}", problems)
    size_bytes = None
    if "size_bytes" in attachment:
        size = attachment.get("size_bytes")
        if isinstance(size, bool) or not isinstance(size, int) or size < 0:
            problems.append(f"{label} size_bytes: expected a non-negative integer, got {_type_name(size)}.")
        else:
            size_bytes = size
    field = evidence_fields.get(raw_field.lower()) if raw_field else None
    if raw_field and field is None:
        problems.append(
            f"{label} u_intake_field: {raw_field} is not an evidence choice. "
            "Assessment answers and the trust-center URL belong on the ticket.")
    return _attachment_preview(file_name, field or raw_field, option, content_type, size_bytes, sys_id), field, option


def _parse_ticket(ticket, row_number, as_of, evidence_fields):
    row = _blank_row(row_number)
    if not isinstance(ticket, dict):
        row["problems"].append(f"Ticket {row_number}: expected an object.")
        return row
    body = ticket_body_columns()
    required = {"number", "category", "attachments", *[f"u_{column}" for column in body]}
    optional = set(OPTIONAL_CONTEXT)
    unknown = sorted(set(ticket) - required - optional)
    shape_problems = []
    if unknown:
        shape_problems.append(
            f"Ticket {row_number}: unrecognized fields {', '.join(unknown)}. They were not dropped.")
    missing = sorted(key for key in required if key not in ticket)
    if missing:
        shape_problems.append(f"Ticket {row_number}: missing {', '.join(missing)}.")
    form = _empty_form()
    for column in body:
        key = f"u_{column}"
        if key in ticket:
            form[column] = _json_text(ticket.get(key), f"Ticket {row_number} {key}", shape_problems)
    number = ""
    if "number" in ticket:
        number = _json_text(ticket.get("number"), f"Ticket {row_number} number", shape_problems)
    if not number and "number" in ticket and not any("number:" in problem for problem in shape_problems):
        shape_problems.append(f"Ticket {row_number}: ticket number is required.")
    if number:
        row["ticket_number"] = number
    if "category" in ticket:
        category = _json_text(ticket.get("category"), f"Ticket {row_number} category", shape_problems)
        if category and category != SERVICENOW_CATEGORY:
            shape_problems.append(f"Ticket {row_number}, category: expected {SERVICENOW_CATEGORY}.")
        if "category" in ticket and category == "" and not any("category:" in item for item in shape_problems):
            shape_problems.append(f"Ticket {row_number}, category: expected {SERVICENOW_CATEGORY}.")
    context = {}
    for key in OPTIONAL_CONTEXT:
        if key not in ticket:
            continue
        text = _json_text(ticket.get(key), f"Ticket {row_number} {key}", shape_problems)
        if text:
            context[key] = text
    row["ticket_context"] = context

    attachments = ticket.get("attachments")
    parsed_attachments = []
    seen = {}
    if "attachments" in ticket and not isinstance(attachments, list):
        shape_problems.append(f"Ticket {row_number}: attachments must be a list.")
    elif isinstance(attachments, list):
        if len(attachments) > MAX_ATTACHMENTS:
            shape_problems.append(
                f"Ticket {row_number} has {len(attachments)} attachments. "
                f"The limit is {MAX_ATTACHMENTS}, so this ticket was not imported.")
        else:
            for index, attachment in enumerate(attachments, start=1):
                label = f"Ticket {row_number} attachment {index}"
                parsed = _parse_attachment(attachment, label, evidence_fields, shape_problems)
                if parsed is None:
                    continue
                preview, field, option = parsed
                parsed_attachments.append(preview)
                if field is None:
                    continue
                if field in seen:
                    previous = seen[field]
                    if previous == option:
                        shape_problems.append(
                            f"Ticket {row_number}: more than one attachment for {field}. "
                            "The extra record was not dropped.")
                    else:
                        shape_problems.append(
                            f"Ticket {row_number}: attachments for {field} disagree "
                            f"({previous or 'blank'} vs {option or 'blank'}).")
                    form[field] = ""
                else:
                    seen[field] = option
                    form[field] = option
            for field in evidence_column_names():
                if field not in seen:
                    shape_problems.append(f"Ticket {row_number}: no attachment metadata for {field}.")
    row["attachments"] = parsed_attachments
    if shape_problems:
        row["problems"].extend(shape_problems)
        row["form"] = form
        row["vendor_name"] = form.get("vendor_name", "")
        return row
    _remember_intake(row, form, as_of, [])
    return row


def _parse_servicenow_json(data, filename, as_of):
    if data[:2] == b"PK":
        return _preview(None, filename, ["The file name says JSON but the contents are a workbook."])
    text, failed = _read_text(data, filename)
    if failed:
        return failed
    if not text.strip():
        return _preview(None, filename, ["The file is empty."])
    try:
        document = json.loads(text, object_pairs_hook=_pairs)
    except _DuplicateKey as exc:
        return _preview(None, filename, [f"Duplicate JSON key {exc.key!r}. The repeated value was not dropped."])
    except json.JSONDecodeError:
        return _preview(None, filename, ["The file is not valid JSON."])
    exported = _reject_known_export(document, filename)
    if exported:
        return exported
    if isinstance(document, list):
        return _preview(None, filename, ["A ServiceNow intake file must be an object with a result list, not a bare array."])
    if not isinstance(document, dict):
        return _preview(None, filename, ["A ServiceNow intake file must be a JSON object."])
    problems = []
    unknown = sorted(set(document) - set(ROOT_KEYS))
    if unknown:
        problems.append("Unrecognized fields: " + ", ".join(unknown) + ". They were not dropped.")
    kind = document.get("import_kind")
    if kind != SERVICENOW_KIND:
        problems.append(f"import_kind must be {SERVICENOW_KIND}. This file is not read from a live ServiceNow instance.")
    if document.get("simulation") is not True:
        problems.append("simulation must be true. This prototype does not call ServiceNow.")
    note = document.get("note")
    if not isinstance(note, str) or not note.strip():
        problems.append("note must be text stating that the file is synthetic and was not retrieved from ServiceNow.")
    elif len(note) > MAX_CELL_LENGTH:
        problems.append(f"note is longer than {MAX_CELL_LENGTH} characters.")
    result = document.get("result")
    if not isinstance(result, list):
        problems.append("result must be a list of tickets.")
        return _preview("servicenow_json", filename, problems)
    if not result:
        problems.append("The file has no vendor rows.")
        return _preview("servicenow_json", filename, problems)
    if len(result) > MAX_DATA_ROWS:
        problems.append(_too_many(len(result)))
        return _preview("servicenow_json", filename, problems)
    if problems:
        return _preview("servicenow_json", filename, problems)
    evidence_fields = _canonical_map(evidence_column_names())
    rows = [_parse_ticket(ticket, index, as_of, evidence_fields) for index, ticket in enumerate(result, start=1)]
    return _finish("servicenow_json", filename, [], rows, 0)


def parse_upload(filename, data, as_of):
    """Parse one upload into a preview. Validation problems stay on the preview."""
    safe = _safe_filename(filename)
    if not isinstance(data, (bytes, bytearray)):
        return _preview(None, safe, ["The upload was not a file."])
    data = bytes(data)
    if len(data) > IMPORT_MAX_BYTES:
        return _preview(None, safe, [f"The file is larger than {IMPORT_MAX_BYTES // 1024} KB."])
    if not data:
        return _preview(None, safe, ["The file is empty."])
    suffix = Path(safe).suffix.lower()
    if suffix == ".xls" or suffix == ".xlsm":
        return _preview(None, safe, ["Only .xlsx workbooks are accepted. Older Excel files and macro files are not."])
    if suffix == ".csv":
        if data[:2] == b"PK":
            return _preview(None, safe, ["The file name says CSV but the contents are a workbook."])
        return _parse_csv(data, safe, as_of)
    if suffix == ".xlsx":
        return _parse_xlsx(data, safe, as_of)
    if suffix == ".json":
        return _parse_servicenow_json(data, safe, as_of)
    return _preview(None, safe, ["The file needs a .csv, .xlsx, or .json extension."])


def _preview_from_confirmation(document, as_of):
    if not isinstance(document, dict):
        return _preview(None, "confirmation", ["The confirmation was not a JSON object."])
    unknown = sorted(set(document) - {"format", "rows"})
    problems = []
    if unknown:
        problems.append("Unrecognized confirmation fields: " + ", ".join(unknown) + ".")
    fmt = document.get("format")
    if fmt not in FORMAT_LABELS or fmt == "unrecognized":
        problems.append("The confirmation format is not an intake import.")
        fmt = "unrecognized"
    rows_in = document.get("rows")
    if not isinstance(rows_in, list) or not rows_in:
        problems.append("The confirmation has no vendor rows.")
        return _preview(fmt, "confirmation", problems)
    if len(rows_in) > MAX_DATA_ROWS:
        problems.append(_too_many(len(rows_in)))
        return _preview(fmt, "confirmation", problems)
    columns = questionnaire_columns()
    rows = []
    for index, item in enumerate(rows_in, start=1):
        row = _blank_row(index)
        if not isinstance(item, dict):
            row["problems"].append(f"Row {index}: expected an object.")
            rows.append(row)
            continue
        extra = sorted(set(item) - {"row_number", "ticket_number", "form"})
        if extra:
            row["problems"].append(f"Row {index}: unrecognized confirmation fields {', '.join(extra)}.")
        number = item.get("row_number")
        if isinstance(number, bool) or not isinstance(number, int) or number < 1:
            row["problems"].append(f"Row {index}: row_number must be a positive integer.")
        else:
            row["row_number"] = number
        ticket = item.get("ticket_number")
        servicenow = fmt.startswith("servicenow")
        if servicenow:
            if not isinstance(ticket, str) or not ticket.strip():
                row["problems"].append(f"Row {row['row_number']}: ticket number is required.")
            else:
                row["ticket_number"] = ticket.strip()
        elif ticket not in (None, ""):
            row["problems"].append(f"Row {row['row_number']}: a questionnaire row has no ServiceNow ticket number.")
        form_in = item.get("form")
        form = _empty_form()
        if not isinstance(form_in, dict):
            row["problems"].append(f"Row {row['row_number']}: form must be an object.")
            rows.append(row)
            continue
        unknown_fields = sorted(set(form_in) - set(columns))
        if unknown_fields:
            row["problems"].append(
                f"Row {row['row_number']}: unrecognized fields {', '.join(unknown_fields)}. They were not dropped.")
        missing_fields = [column for column in columns if column not in form_in]
        if missing_fields:
            row["problems"].append(f"Row {row['row_number']}: missing {', '.join(missing_fields)}.")
        shape = list(row["problems"])
        for column in columns:
            if column not in form_in:
                continue
            value = form_in[column]
            if not isinstance(value, str):
                shape.append(f"Row {row['row_number']}, {column}: expected text, got {_type_name(value)}.")
                continue
            text, problem = _cell_text(value)
            if problem:
                shape.append(f"Row {row['row_number']}, {column}: {problem}.")
            else:
                form[column] = text
        if shape:
            row["problems"] = shape
            row["form"] = form
            row["vendor_name"] = form.get("vendor_name", "")
        else:
            _remember_intake(row, form, as_of, [])
        rows.append(row)
    if problems:
        return _preview(fmt, "confirmation", problems)
    return _finish(fmt, "confirmation", [], rows, 0)


def _import_summary(fmt, row):
    label = FORMAT_LABELS.get(fmt, "file")
    text = f"Imported from row {row['row_number']} of a {label} file. Synthetic input, not saved."
    if row.get("ticket_number"):
        text += (f" Ticket {row['ticket_number']} is simulated file context "
                 "and was not read from ServiceNow.")
    return text


def assess_import(document, as_of):
    """Re-check a confirmation document and run each row through intake."""
    preview = _preview_from_confirmation(document, as_of)
    if not can_import(preview):
        raise ImportRejected(preview)
    results = []
    for row in preview["rows"]:
        intake = build_intake(row["form"], as_of)
        result = run_intake(intake, as_of)
        result["summary"] = _import_summary(preview["format"], row)
        result["source_row"] = row["row_number"]
        result["ticket_number"] = row["ticket_number"]
        result["import_format"] = preview["format"]
        results.append(result)
    return results


def sample_forms():
    """Two synthetic vendors used by the downloadable samples and the fixtures."""
    questionnaire = load_questionnaire()
    policy = load_org_policy()
    north = _empty_form()
    north.update({
        "vendor_name": "Northwind Notes",
        "data_sensitivity": "internal",
        "integration_permissions": "read_only",
        "ai_use": "none",
        "model_training_on_customer_data": "no",
        "external_exposure": "internal_only",
        "retention_deletion": "documented",
    })
    for spec in policy["intake_fields"]:
        north[spec["field"]] = "no"
    north[questionnaire["access_evidence"]["field"]] = "independent_verified"
    north[questionnaire["trust_center"]["url_field"]] = "https://trust.example.test/northwind"
    for pointer in questionnaire["trust_center"]["pointers"]:
        north[pointer["field"]] = "independent_verified"
    for clause in policy["required_contract_clauses"]:
        north[clause_evidence_field(clause["clause_id"])] = "independent_verified"

    harbor = _empty_form()
    harbor.update({
        "vendor_name": "Harbor Assist",
        "data_sensitivity": "confidential",
        "integration_permissions": "read_write",
        "ai_use": "assistive",
        "model_training_on_customer_data": "no",
        "external_exposure": "customer_facing",
        "retention_deletion": "documented",
    })
    for spec in policy["intake_fields"]:
        harbor[spec["field"]] = "no"
    harbor[questionnaire["access_evidence"]["field"]] = "vendor_statement"
    harbor[questionnaire["trust_center"]["url_field"]] = ""
    harbor_pointers = {"soc2": "vendor_statement", "security": "not_reviewed", "privacy": "none"}
    for pointer in questionnaire["trust_center"]["pointers"]:
        harbor[pointer["field"]] = harbor_pointers.get(pointer["pointer_id"], "none")
    for clause in policy["required_contract_clauses"]:
        harbor[clause_evidence_field(clause["clause_id"])] = "vendor_statement"
    return [north, harbor]


def _ticket_context(form, index):
    return {
        "sys_id": f"synthetic-ticket-{index:04d}",
        "number": f"VEND001000{index}",
        "short_description": f"Vendor intake: {form['vendor_name']}",
        "description": "Synthetic vendor intake ticket for the portfolio demo.",
        "category": SERVICENOW_CATEGORY,
        "state": "2",
        "priority": "4",
        "assignment_group": "GRC Review (demo)",
        "opened_at": "2026-10-01 09:00:00",
    }


def _servicenow_document(forms):
    evidence = evidence_column_names()
    result = []
    for index, form in enumerate(forms, start=1):
        ticket = _ticket_context(form, index)
        for column in ticket_body_columns():
            ticket[f"u_{column}"] = form[column]
        ticket["attachments"] = [
            {
                "sys_id": f"synthetic-att-{index:04d}-{field}",
                "file_name": f"{field}.txt",
                "content_type": "text/plain",
                "size_bytes": 64,
                "table_name": "incident",
                "table_sys_id": ticket["sys_id"],
                "u_intake_field": field,
                "u_evidence_option": form[field],
            }
            for field in evidence
        ]
        result.append(ticket)
    return {
        "import_kind": SERVICENOW_KIND,
        "simulation": True,
        "note": "Synthetic file. Not retrieved from a live ServiceNow instance.",
        "result": result,
    }


def render_questionnaire_csv(forms=None):
    forms = sample_forms() if forms is None else forms
    buffer = io.StringIO(newline="")
    columns = questionnaire_columns()
    writer = csv.DictWriter(buffer, fieldnames=columns, lineterminator="\n", extrasaction="raise")
    writer.writeheader()
    for form in forms:
        writer.writerow({column: form.get(column, "") for column in columns})
    return buffer.getvalue().encode("utf-8")


def _freeze_xlsx(data):
    """Rewrite a workbook so sample bytes do not change with the clock."""
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = re.sub(r"\d{4}-\d{2}-\d{2}T[^<]+", "2026-10-23T00:00:00Z", payload.decode("utf-8"))
                payload = text.encode("utf-8")
            fixed = zipfile.ZipInfo(filename=info.filename, date_time=(2026, 10, 23, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            target.writestr(fixed, payload)
    source.close()
    return buffer.getvalue()


def render_questionnaire_xlsx(forms=None):
    from openpyxl import Workbook

    forms = sample_forms() if forms is None else forms
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Intake"
    columns = questionnaire_columns()
    sheet.append(columns)
    for form in forms:
        sheet.append([form.get(column, "") for column in columns])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return _freeze_xlsx(buffer.getvalue())


def render_servicenow_json(forms=None):
    forms = sample_forms() if forms is None else forms
    return (json.dumps(_servicenow_document(forms), indent=2) + "\n").encode("utf-8")


def render_servicenow_csv(forms=None):
    forms = sample_forms() if forms is None else forms
    columns = ["number", "short_description", "description", "category", "state", "priority",
               "assignment_group", "opened_at", "sys_id", *[f"u_{column}" for column in questionnaire_columns()]]
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=columns, lineterminator="\n")
    writer.writeheader()
    for index, form in enumerate(forms, start=1):
        context = _ticket_context(form, index)
        row = {key: context[key] for key in (
            "number", "short_description", "description", "category", "state", "priority",
            "assignment_group", "opened_at", "sys_id")}
        for column in questionnaire_columns():
            row[f"u_{column}"] = form[column]
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8")


def sample_file(kind):
    if kind == "csv":
        return render_questionnaire_csv()
    if kind == "xlsx":
        return render_questionnaire_xlsx()
    if kind == "json":
        return render_servicenow_json()
    if kind == "servicenow-csv":
        return render_servicenow_csv()
    raise ValueError(kind)
