"""File import: questionnaire sheets and simulated ServiceNow tickets."""

import base64
import io
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment

from app import create_app
from app.services.intake_builder import build_intake, clause_evidence_field, remediation_payload, run_intake
from app.services.intake_export import build_package, render_json
from app.services.intake_importer import (
    IMPORT_MAX_BYTES, ImportRejected, assess_import, can_import, confirmation_document,
    parse_upload, questionnaire_columns, render_questionnaire_csv, sample_file, sample_forms,
)
from app.services.org_policy import load_org_policy

AS_OF = date(2026, 10, 23)
FIXTURES = Path(__file__).parent / "fixtures" / "imports"
SAME_ORIGIN = {"Origin": "http://localhost"}
FIXTURE_KINDS = {
    "vendor_intake.csv": "csv",
    "vendor_intake.xlsx": "xlsx",
    "servicenow_intake.json": "json",
    "servicenow_intake.csv": "servicenow-csv",
}


def _load(name):
    return (FIXTURES / name).read_bytes()


def _forms(preview):
    return [row["form"] for row in preview["rows"]]


def _assess(preview):
    return assess_import(confirmation_document(preview), AS_OF)


def _direct(form):
    return run_intake(build_intake(form, AS_OF), AS_OF)


@pytest.fixture
def gated_client(tmp_path):
    app = create_app({"TESTING": True, "DATABASE": str(tmp_path / "import.sqlite3"),
                      "REQUIRE_AUTH": True, "DEMO_USERNAME": "demo", "DEMO_PASSWORD": "pw"})
    return app.test_client()


def test_fixtures_match_the_sample_generator():
    for name, kind in FIXTURE_KINDS.items():
        assert _load(name) == sample_file(kind)


@pytest.mark.parametrize("name", ["vendor_intake.csv", "vendor_intake.xlsx", "servicenow_intake.json", "servicenow_intake.csv"])
def test_happy_path_fixtures_match_the_questionnaire(name):
    preview = parse_upload(name, _load(name), AS_OF)
    assert can_import(preview)
    assert preview["problems"] == []
    assert _forms(preview) == sample_forms()
    imported = _assess(preview)
    direct = [_direct(form) for form in sample_forms()]
    assert [item["risk"] for item in imported] == [item["risk"] for item in direct]
    assert imported[0]["risk"]["inherent"]["score"] == 20
    assert imported[0]["risk"]["inherent"]["tier"] == "low"
    assert imported[0]["trust_center"]["url"] == "https://trust.example.test/northwind"
    assert imported[0]["trust_center"]["fetched"] is False
    assert "not saved" in imported[0]["summary"]
    assert imported[0]["risk"]["policy_outcome"]["status"] == "not_prohibited"


def test_servicenow_ticket_context_is_not_scored():
    document = json.loads(sample_file("json"))
    document["result"][0]["priority"] = "1"
    document["result"][0]["opened_at"] = "1999-01-01 00:00:00"
    document["result"][0]["state"] = "7"
    preview = parse_upload("tickets.json", json.dumps(document).encode(), AS_OF)
    assert can_import(preview)
    assert preview["rows"][0]["ticket_number"] == "VEND0010001"
    assert preview["rows"][0]["ticket_context"]["priority"] == "1"
    imported = _assess(preview)[0]
    assert imported["risk"] == _direct(sample_forms()[0])["risk"]
    assert imported["as_of"] == AS_OF.isoformat()
    assert "was not read from ServiceNow" in imported["summary"]


def test_missing_columns_are_reported_and_block_import():
    forms = sample_forms()
    kept = [column for column in questionnaire_columns() if column != "data_sensitivity"]
    lines = [",".join(kept)]
    for form in forms:
        lines.append(",".join(form[column] for column in kept))
    preview = parse_upload("missing.csv", ("\n".join(lines) + "\n").encode(), AS_OF)
    assert not can_import(preview)
    assert any("Missing columns: data_sensitivity" in problem for problem in preview["problems"])
    assert preview["rows"] == []


def test_unknown_column_is_not_dropped():
    forms = sample_forms()[:1]
    columns = [*questionnaire_columns(), "reviewer_notes"]
    line = ",".join(columns)
    values = ",".join([forms[0][column] for column in questionnaire_columns()] + ["keep me"])
    preview = parse_upload("notes.csv", f"{line}\n{values}\n".encode(), AS_OF)
    assert not can_import(preview)
    assert any("reviewer_notes" in problem and "not dropped" in problem for problem in preview["problems"])


def test_bad_types_in_xlsx_are_not_coerced():
    columns = questionnaire_columns()
    index = columns.index("data_sensitivity")

    def workbook(value):
        book = Workbook()
        sheet = book.active
        sheet.append(columns)
        row = [sample_forms()[0][column] for column in columns]
        row[index] = value
        sheet.append(row)
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        return buffer.getvalue()

    number = parse_upload("numbers.xlsx", workbook(10), AS_OF)
    assert not can_import(number)
    assert any("expected text, got a number" in problem for problem in number["rows"][0]["problems"])
    boolean = parse_upload("flags.xlsx", workbook(True), AS_OF)
    assert any("expected text, got a boolean" in problem for problem in boolean["rows"][0]["problems"])
    formula = parse_upload("formula.xlsx", workbook("=1+1"), AS_OF)
    assert any("formulas are not accepted" in problem for problem in formula["rows"][0]["problems"])


def test_bad_json_types_are_not_coerced():
    document = json.loads(sample_file("json"))
    document["result"][0]["u_data_sensitivity"] = 10
    preview = parse_upload("bad.json", json.dumps(document).encode(), AS_OF)
    assert not can_import(preview)
    assert any("u_data_sensitivity" in problem and "a number" in problem for problem in preview["rows"][0]["problems"])

    document = json.loads(sample_file("json"))
    document["result"][0]["u_ai_use"] = None
    preview = parse_upload("null.json", json.dumps(document).encode(), AS_OF)
    assert any("u_ai_use" in problem and "null" in problem for problem in preview["rows"][0]["problems"])


@pytest.mark.parametrize("filename, payload", [
    ("empty.csv", b""),
    ("blank.csv", b"   \n"),
    ("header.csv", (",".join(questionnaire_columns()) + "\n").encode()),
    ("empty.json", b""),
    ("no-rows.json", json.dumps({
        "import_kind": "servicenow_vendor_intake",
        "simulation": True,
        "note": "Synthetic file. Not retrieved from a live ServiceNow instance.",
        "result": [],
    }).encode()),
])
def test_empty_files_are_rejected(filename, payload):
    preview = parse_upload(filename, payload, AS_OF)
    assert not can_import(preview)
    assert preview["rows"] == []
    assert preview["problems"]


def test_empty_workbook_is_rejected():
    book = Workbook()
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    preview = parse_upload("empty.xlsx", buffer.getvalue(), AS_OF)
    assert not can_import(preview)
    assert any("header" in problem.lower() or "no vendor" in problem.lower() for problem in preview["problems"])


def test_conflicting_rows_are_kept_and_not_imported():
    forms = [dict(sample_forms()[0]), dict(sample_forms()[0])]
    forms[1]["data_sensitivity"] = "restricted"
    preview = parse_upload("conflict.csv", render_questionnaire_csv(forms), AS_OF)
    assert not can_import(preview)
    assert len(preview["rows"]) == 2
    assert preview["rows"][0]["conflicts"]
    assert "disagree" in preview["rows"][0]["conflicts"][0]
    assert "data_sensitivity" in preview["rows"][0]["conflicts"][0]
    assert "restricted" in preview["rows"][0]["conflicts"][0]
    with pytest.raises(ImportRejected):
        _assess(preview)


def test_duplicate_rows_are_not_silently_dropped():
    forms = [sample_forms()[0], sample_forms()[0]]
    preview = parse_upload("dupes.csv", render_questionnaire_csv(forms), AS_OF)
    assert len(preview["rows"]) == 2
    assert not can_import(preview)
    assert "same answers" in preview["rows"][1]["conflicts"][0]


def test_attachment_conflicts_are_reported():
    document = json.loads(sample_file("json"))
    first = document["result"][0]
    soc = next(item for item in first["attachments"] if item["u_intake_field"] == "trust_soc2")
    first["attachments"].append({**soc, "sys_id": "synthetic-att-extra", "u_evidence_option": "contradicts"})
    preview = parse_upload("conflict.json", json.dumps(document).encode(), AS_OF)
    assert not can_import(preview)
    assert any("trust_soc2" in problem and "disagree" in problem for problem in preview["rows"][0]["problems"])
    assert len([item for item in preview["rows"][0]["attachments"] if item["intake_field"] == "trust_soc2"]) == 2


def test_same_vendor_on_two_tickets_conflicts():
    document = json.loads(sample_file("json"))
    document["result"][1]["u_vendor_name"] = document["result"][0]["u_vendor_name"]
    preview = parse_upload("tickets.json", json.dumps(document).encode(), AS_OF)
    assert not can_import(preview)
    assert preview["rows"][0]["conflicts"]


def test_blank_training_is_not_applicable_only_when_ai_use_is_none():
    form = dict(sample_forms()[0])
    form["model_training_on_customer_data"] = ""
    preview = parse_upload("training.csv", render_questionnaire_csv([form]), AS_OF)
    assert can_import(preview)
    result = _assess(preview)[0]
    assert result["answer_applicability"]["model_training_on_customer_data"] == "not_applicable"
    assert result["risk"]["inherent"]["score"] == 20

    form["ai_use"] = "assistive"
    preview = parse_upload("training.csv", render_questionnaire_csv([form]), AS_OF)
    assert not can_import(preview)
    assert any("train or improve models" in problem for problem in preview["rows"][0]["problems"])


def test_extra_worksheet_values_are_not_dropped():
    book = Workbook()
    sheet = book.active
    sheet.append(questionnaire_columns())
    sheet.append([sample_forms()[0][column] for column in questionnaire_columns()])
    other = book.create_sheet("Notes")
    other["A1"] = "keep this note"
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    preview = parse_upload("extra.xlsx", buffer.getvalue(), AS_OF)
    assert not can_import(preview)
    assert any("Notes" in problem for problem in preview["problems"])


def test_merged_cells_and_comments_block_import():
    columns = questionnaire_columns()
    book = Workbook()
    sheet = book.active
    sheet.append(columns)
    sheet.append([sample_forms()[0][column] for column in columns])
    sheet.merge_cells("A1:B1")
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    merged = parse_upload("merged.xlsx", buffer.getvalue(), AS_OF)
    assert any("Merged cells" in problem for problem in merged["problems"])

    book = Workbook()
    sheet = book.active
    sheet.append(columns)
    sheet.append([sample_forms()[0][column] for column in columns])
    sheet.cell(2, 1).comment = Comment("not evidence", "Demo")
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    commented = parse_upload("comment.xlsx", buffer.getvalue(), AS_OF)
    assert any("comments" in problem for problem in commented["problems"])
    assert not can_import(commented)


def test_reviewer_package_and_remediation_ticket_are_not_intake_files():
    package = render_json(build_package(_direct(sample_forms()[0])))
    preview = parse_upload("package.json", package.encode(), AS_OF)
    assert any("reviewer package" in problem for problem in preview["problems"])

    payload = remediation_payload(_direct({**sample_forms()[1]}))
    preview = parse_upload("ticket.json", json.dumps(payload).encode(), AS_OF)
    assert any("remediation ticket" in problem for problem in preview["problems"])


def test_simulation_flag_is_required():
    document = json.loads(sample_file("json"))
    document["simulation"] = False
    preview = parse_upload("live.json", json.dumps(document).encode(), AS_OF)
    assert any("does not call ServiceNow" in problem for problem in preview["problems"])
    assert preview["rows"] == []


def test_semicolon_csv_is_rejected():
    preview = parse_upload("eu.csv", b"vendor_name;data_sensitivity\nNorthwind;internal\n", AS_OF)
    assert any("commas" in problem for problem in preview["problems"])


def test_duplicate_json_key_is_not_dropped():
    text = sample_file("json").decode().replace('"priority": "4"', '"priority": "4", "priority": "1"', 1)
    preview = parse_upload("dupes.json", text.encode(), AS_OF)
    assert any("Duplicate JSON key 'priority'" in problem for problem in preview["problems"])


def test_wrong_category_blocks_the_ticket():
    document = json.loads(sample_file("json"))
    document["result"][0]["category"] = "hardware"
    preview = parse_upload("hardware.json", json.dumps(document).encode(), AS_OF)
    assert not can_import(preview)
    assert any("vendor_risk" in problem for problem in preview["rows"][0]["problems"])


def test_clause_column_is_the_live_policy_field():
    clause_field = clause_evidence_field(load_org_policy()["required_contract_clauses"][0]["clause_id"])
    assert clause_field in questionnaire_columns()
    assert sample_forms()[0][clause_field] == "independent_verified"


def _post_file(client, name, payload):
    return client.post("/try/import/preview", data={"intake_file": (io.BytesIO(payload), name)},
                       headers=SAME_ORIGIN, content_type="multipart/form-data")


def test_import_page_explains_the_simulation(gated_client):
    response = gated_client.get("/try/import")
    assert response.status_code == 200
    assert b"not a live connector" in response.data
    assert b"servicenow_vendor_intake" in response.data
    assert b"data_sensitivity" in response.data
    page = gated_client.get("/try/")
    assert b"/try/import" in page.data


def test_preview_does_not_show_a_score_and_confirm_does(gated_client):
    response = _post_file(gated_client, "vendor_intake.csv", sample_file("csv"))
    assert response.status_code == 200
    assert b"Northwind Notes" in response.data
    assert b"Harbor Assist" in response.data
    assert b"Assess these vendors" in response.data
    assert b"INHERENT RISK" not in response.data
    start = response.data.index(b"<textarea")
    snippet = response.data[start:response.data.index(b"</textarea>", start)]
    document = json.loads(snippet.split(b">", 1)[1])
    confirmed = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                  headers=SAME_ORIGIN)
    assert confirmed.status_code == 200
    assert b"INHERENT RISK" in confirmed.data
    assert b"Northwind Notes" in confirmed.data
    assert b"not saved" in confirmed.data
    assert b"was not read from ServiceNow" not in confirmed.data
    json_preview = _post_file(gated_client, "servicenow_intake.json", sample_file("json"))
    assert json_preview.status_code == 200
    assert b"trust_soc2.txt" in json_preview.data
    assert b"VEND0010001" in json_preview.data
    assert b"INHERENT RISK" not in json_preview.data
    assert b"do not set residual risk" in json_preview.data
    login = {"Authorization": "Basic " + base64.b64encode(b"demo:pw").decode()}
    register = gated_client.get("/risks", headers=login)
    assert b"Northwind Notes" not in register.data
    assert b"Harbor Assist" not in register.data


def _confirmation(page):
    start = page.index(b"<textarea")
    snippet = page[start:page.index(b"</textarea>", start)]
    return json.loads(snippet.split(b">", 1)[1])


def test_import_save_is_opt_in_and_a_low_tier_is_not_approval(gated_client):
    response = _post_file(gated_client, "vendor_intake.csv", sample_file("csv"))
    document = _confirmation(response.data)
    confirmed = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                  headers=SAME_ORIGIN)
    assert confirmed.status_code == 200
    assert b"not saved" in confirmed.data
    assert b"Save this synthetic vendor" in confirmed.data
    login = {"Authorization": "Basic " + base64.b64encode(b"demo:pw").decode()}
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    row = document["rows"][0]
    fields = {
        "confirmation": json.dumps(document),
        "source_row": row["row_number"],
        "disposition": "escalate",
        "rationale": "Low residual tier is not an approval.",
        "decider": "GRC analyst",
    }
    missing_box = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert missing_box.status_code == 400
    assert b"Nothing was saved" in missing_box.data
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    invalid = dict(fields)
    invalid["persist"] = "yes"
    invalid["disposition"] = "accept"
    rejected = gated_client.post("/try/import/save", data=invalid, headers=SAME_ORIGIN)
    assert rejected.status_code == 400
    assert b"disposition" in rejected.data
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    cross_site = dict(fields)
    cross_site["persist"] = "yes"
    blocked = gated_client.post("/try/import/save", data=cross_site)
    assert blocked.status_code == 403

    saved_fields = dict(fields)
    saved_fields["persist"] = "yes"
    saved = gated_client.post("/try/import/save", data=saved_fields, headers=SAME_ORIGIN)
    assert saved.status_code == 302
    detail = gated_client.get(saved.headers["Location"], headers=login)
    assert detail.status_code == 200
    assert b"Northwind Notes" in detail.data
    assert b"Escalate" in detail.data
    assert b"Low" in detail.data
    assert b"not an approval" in detail.data
    register = gated_client.get("/risks", headers=login)
    assert b"Northwind Notes" in register.data
    assert b"Harbor Assist" not in register.data


def test_import_save_stores_optional_minutes_and_rejects_a_negative(gated_client):
    from app.db import get_db
    from app.services.workflow_service import get_risk

    response = _post_file(gated_client, "vendor_intake.csv", sample_file("csv"))
    document = _confirmation(response.data)
    confirmed = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                  headers=SAME_ORIGIN)
    assert b"minutes spent reviewing" in confirmed.data
    assert b'name="handling_minutes" required' not in confirmed.data
    row = document["rows"][0]
    fields = {
        "confirmation": json.dumps(document),
        "source_row": row["row_number"],
        "persist": "yes",
        "disposition": "escalate",
        "rationale": "Low residual tier is not an approval.",
        "decider": "GRC analyst",
        "handling_minutes": "-1",
    }
    rejected = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert rejected.status_code == 400
    assert b"handling_minutes" in rejected.data
    login = {"Authorization": "Basic " + base64.b64encode(b"demo:pw").decode()}
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    fields["handling_minutes"] = "16"
    saved = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert saved.status_code == 302
    risk_id = saved.headers["Location"].rsplit("/", 1)[-1].split("?", 1)[0]
    with gated_client.application.app_context():
        decision = get_risk(get_db(), risk_id)["decision"]
    assert decision["handling_minutes"] == 16
    assert decision["disposition"] == "escalate"
    detail = gated_client.get(saved.headers["Location"], headers=login)
    assert b"16 self-reported minutes" in detail.data
    assert b"Not reported" in detail.data


def test_import_save_stores_optional_rework_and_rejects_a_fraction(gated_client):
    from app.db import get_db
    from app.services.workflow_service import get_risk

    response = _post_file(gated_client, "vendor_intake.csv", sample_file("csv"))
    document = _confirmation(response.data)
    confirmed = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                  headers=SAME_ORIGIN)
    assert b'name="rework_minutes"' in confirmed.data
    assert b'name="rework_minutes" required' not in confirmed.data
    row = document["rows"][0]
    fields = {
        "confirmation": json.dumps(document),
        "source_row": row["row_number"],
        "persist": "yes",
        "disposition": "escalate",
        "rationale": "Low residual tier is not an approval.",
        "decider": "GRC analyst",
        "handling_minutes": "16",
        "rework_minutes": "2.5",
    }
    rejected = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert rejected.status_code == 400
    assert b"rework_minutes" in rejected.data
    login = {"Authorization": "Basic " + base64.b64encode(b"demo:pw").decode()}
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    fields["rework_minutes"] = "8"
    saved = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert saved.status_code == 302
    risk_id = saved.headers["Location"].rsplit("/", 1)[-1].split("?", 1)[0]
    with gated_client.application.app_context():
        decision = get_risk(get_db(), risk_id)["decision"]
    assert decision["handling_minutes"] == 16
    assert decision["rework_minutes"] == 8
    detail = gated_client.get(saved.headers["Location"], headers=login)
    assert b"8 self-reported rework minutes" in detail.data
    assert b"16 self-reported minutes" in detail.data
    assert b"Not reported" in detail.data


def test_import_save_stores_optional_manual_touches_and_rejects_text(gated_client):
    from app.db import get_db
    from app.services.workflow_service import get_risk

    response = _post_file(gated_client, "vendor_intake.csv", sample_file("csv"))
    document = _confirmation(response.data)
    confirmed = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                  headers=SAME_ORIGIN)
    assert b'name="manual_touches"' in confirmed.data
    assert b'name="manual_touches" required' not in confirmed.data
    row = document["rows"][0]
    fields = {
        "confirmation": json.dumps(document),
        "source_row": row["row_number"],
        "persist": "yes",
        "disposition": "escalate",
        "rationale": "Low residual tier is not an approval.",
        "decider": "GRC analyst",
        "handling_minutes": "16",
        "rework_minutes": "8",
        "manual_touches": "many",
    }
    rejected = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert rejected.status_code == 400
    assert b"manual_touches" in rejected.data
    login = {"Authorization": "Basic " + base64.b64encode(b"demo:pw").decode()}
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    fields["manual_touches"] = "3"
    saved = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert saved.status_code == 302
    risk_id = saved.headers["Location"].rsplit("/", 1)[-1].split("?", 1)[0]
    with gated_client.application.app_context():
        decision = get_risk(get_db(), risk_id)["decision"]
    assert decision["handling_minutes"] == 16
    assert decision["rework_minutes"] == 8
    assert decision["manual_touches"] == 3
    detail = gated_client.get(saved.headers["Location"], headers=login)
    assert b"3 self-reported touches" in detail.data
    assert b"8 self-reported rework minutes" in detail.data
    assert b"16 self-reported minutes" in detail.data


def test_import_save_stores_optional_handoffs_and_rejects_a_fraction(gated_client):
    from app.db import get_db
    from app.services.workflow_service import get_risk

    response = _post_file(gated_client, "vendor_intake.csv", sample_file("csv"))
    document = _confirmation(response.data)
    confirmed = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                  headers=SAME_ORIGIN)
    assert b'name="handoffs"' in confirmed.data
    assert b'name="handoffs" required' not in confirmed.data
    assert b'name="authenticated_effort_minutes"' not in confirmed.data
    assert b'name="cash_savings_dollars"' not in confirmed.data
    assert b"does not ask for authenticated effort" in confirmed.data
    row = document["rows"][0]
    fields = {
        "confirmation": json.dumps(document),
        "source_row": row["row_number"],
        "persist": "yes",
        "disposition": "escalate",
        "rationale": "Low residual tier is not an approval.",
        "decider": "GRC analyst",
        "handling_minutes": "16",
        "rework_minutes": "8",
        "manual_touches": "3",
        "handoffs": "1.5",
    }
    rejected = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert rejected.status_code == 400
    assert b"handoffs" in rejected.data
    login = {"Authorization": "Basic " + base64.b64encode(b"demo:pw").decode()}
    assert b"Northwind Notes" not in gated_client.get("/risks", headers=login).data

    fields["handoffs"] = "2"
    saved = gated_client.post("/try/import/save", data=fields, headers=SAME_ORIGIN)
    assert saved.status_code == 302
    risk_id = saved.headers["Location"].rsplit("/", 1)[-1].split("?", 1)[0]
    with gated_client.application.app_context():
        decision = get_risk(get_db(), risk_id)["decision"]
    assert decision["handling_minutes"] == 16
    assert decision["rework_minutes"] == 8
    assert decision["manual_touches"] == 3
    assert decision["handoffs"] == 2
    detail = gated_client.get(saved.headers["Location"], headers=login)
    assert b"2 self-reported handoffs" in detail.data
    assert b"3 self-reported touches" in detail.data
    assert b"8 self-reported rework minutes" in detail.data
    assert b"16 self-reported minutes" in detail.data


def test_preview_shows_row_errors_without_a_confirm_button(gated_client):
    forms = [dict(sample_forms()[0]), dict(sample_forms()[0])]
    forms[1]["data_sensitivity"] = "restricted"
    response = _post_file(gated_client, "conflict.csv", render_questionnaire_csv(forms))
    assert response.status_code == 400
    assert b"disagree" in response.data
    assert b"Assess these vendors" not in response.data
    assert b"INHERENT RISK" not in response.data


def test_sample_downloads_parse(gated_client):
    for kind, name in (("csv", "vendor-intake.csv"), ("xlsx", "vendor-intake.xlsx"),
                       ("json", "servicenow-intake.json"), ("servicenow-csv", "servicenow-intake.csv")):
        response = gated_client.get(f"/try/import/sample/{kind}")
        assert response.status_code == 200
        assert name.encode() in response.headers["Content-Disposition"].encode()
        preview = parse_upload(name, response.data, AS_OF)
        assert can_import(preview)


def test_import_pages_are_public_and_reject_cross_site_posts(gated_client):
    assert gated_client.get("/try/import").status_code == 200
    assert gated_client.get("/").status_code == 401
    uploaded = gated_client.post("/try/import/preview", data={"intake_file": (io.BytesIO(sample_file("csv")), "v.csv")},
                                 headers={"Origin": "https://evil.example"}, content_type="multipart/form-data")
    assert uploaded.status_code == 403


def test_vendor_name_is_escaped_on_preview_and_result(gated_client):
    form = dict(sample_forms()[0])
    form["vendor_name"] = "<script>x</script>"
    response = _post_file(gated_client, "script.csv", render_questionnaire_csv([form]))
    assert response.status_code == 200
    assert b"<script>x</script>" not in response.data
    confirmed = gated_client.post(
        "/try/import/confirm",
        data={"confirmation": json.dumps(confirmation_document(parse_upload("script.csv", render_questionnaire_csv([form]), AS_OF)))},
        headers=SAME_ORIGIN)
    assert confirmed.status_code == 200
    assert b"<script>x</script>" not in confirmed.data


def test_missing_file_is_a_validation_error(gated_client):
    response = gated_client.post("/try/import/preview", data={}, headers=SAME_ORIGIN)
    assert response.status_code == 400
    assert b"Choose a CSV" in response.data


def test_oversized_import_is_rejected_and_small_import_is_not(gated_client):
    too_big = _post_file(gated_client, "big.csv", b"a" * (IMPORT_MAX_BYTES + 50))
    assert too_big.status_code == 413
    wide = sample_file("csv") + (b"\n" + b"x" * 20000)
    accepted = _post_file(gated_client, "wide.csv", wide)
    assert accepted.status_code != 413


def test_unknown_sample_is_not_found(gated_client):
    assert gated_client.get("/try/import/sample/exe").status_code == 404


def test_confirm_rejects_a_tampered_score_field(gated_client):
    document = confirmation_document(parse_upload("vendor_intake.csv", sample_file("csv"), AS_OF))
    document["rows"][0]["form"]["inherent_tier"] = "low"
    response = gated_client.post("/try/import/confirm", data={"confirmation": json.dumps(document)},
                                 headers=SAME_ORIGIN)
    assert response.status_code == 400
    assert b"inherent_tier" in response.data
    assert b"not dropped" in response.data
    assert b"INHERENT RISK" not in response.data
